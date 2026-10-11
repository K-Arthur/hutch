"""PPTX generator: E8 autofit defense, geometric gate, bounds rejection, preview shape."""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

import pptx as pptx_library

import hutch.documents.docgen.pptx as pptx_gen  # registers the "pptx" format
from hutch.documents.docgen import DocumentGenerationError, generate

_HAS_PROBED_FONT = pptx_gen.probe_font() is not None


def _spec() -> dict:
    """One deck exercising every layout, bullets, a table, and notes."""
    return {"slides": [
        {"layout": "title", "title": "Quarterly Review", "subtitle": "FY 2026",
         "notes": "welcome everyone\nsecond line"},
        {"title": "Roadmap", "bullets": ["ship docgen", "harden gates"],
         "table": {"header": ["area", "state"], "rows": [["pptx", "verified"]]}},
        {"layout": "section", "title": "Questions", "subtitle": "ask away"},
    ]}


def _named_shapes(slide):
    return {shape.name: shape for shape in slide.shapes}


def test_happy_path_roundtrip_preview_and_generator_string():
    result = generate("pptx", _spec())
    assert result.generator == \
        f"hutch docgen pptx 1.0 (python-pptx {pptx_library.__version__})"
    assert result.validation_state == "passed" and result.findings == []
    assert result.preview_mode == "structure"
    assert result.preview_data == {"kind": "slides", "items": [
        {"title": "Quarterly Review", "bullet_count": 0, "has_table": False, "has_notes": True},
        {"title": "Roadmap", "bullet_count": 2, "has_table": True, "has_notes": False},
        {"title": "Questions", "bullet_count": 0, "has_table": False, "has_notes": False}]}
    assert result.page_count == 3

    deck = Presentation(io.BytesIO(result.data))  # parse-back of the exact stored bytes
    assert len(deck.slides) == 3
    first = _named_shapes(deck.slides[0])
    assert first["hutch-title"].text_frame.text == "Quarterly Review"
    assert first["hutch-subtitle"].text_frame.text == "FY 2026"
    assert deck.slides[0].has_notes_slide
    assert deck.slides[0].notes_slide.notes_text_frame.text == "welcome everyone\nsecond line"

    second = _named_shapes(deck.slides[1])
    bullets = second["hutch-bullets"].text_frame
    assert [p.text for p in bullets.paragraphs] == ["ship docgen", "harden gates"]
    table = second["hutch-table"].table
    assert len(table.columns) == 2 and len(table.rows) == 2
    assert table.cell(0, 0).text == "area" and table.cell(1, 1).text == "verified"
    assert not deck.slides[1].has_notes_slide

    third = _named_shapes(deck.slides[2])
    assert third["hutch-title"].text_frame.text == "Questions"
    assert third["hutch-subtitle"].text_frame.text == "ask away"

    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height


def test_fit_bakes_explicit_sizes_instead_of_renderer_autofit():
    """python-pptx#973: autofit is renderer metadata; the bytes must already fit (E8)."""
    result = generate("pptx", {"slides": [{"title": "T", "bullets": ["short one", "another"]}]})
    frame = _named_shapes(Presentation(io.BytesIO(result.data)).slides[0])["hutch-bullets"]
    sizes = [run.font.size for p in frame.text_frame.paragraphs for run in p.runs]
    assert sizes and all(size is not None for size in sizes), \
        "run sizes must be baked into the file, never deferred to the viewer"
    if _HAS_PROBED_FONT:
        assert result.validation_state == "passed" and result.findings == []
        assert frame.text_frame.auto_size is MSO_AUTO_SIZE.NONE, \
            "fit_text must disable autofit, not write TEXT_TO_FIT_SHAPE"
    else:
        assert [f["code"] for f in result.findings] == ["fit_text_unavailable"]
        assert result.validation_state == "passed_with_warnings"


def test_unbreakable_word_falls_back_to_fixed_sizes():
    """A word wider than its frame at any fit size crashes TextFitter; the deck must
    still be produced with conservative sizes and a disclosed finding."""
    result = generate("pptx", {"slides": [
        {"title": "T", "bullets": ["see " + "x" * 150 + " for details"]}]})
    if _HAS_PROBED_FONT:
        assert "fit_text_fallback" in [f["code"] for f in result.findings]
        assert result.validation_state == "passed_with_warnings"
    deck = Presentation(io.BytesIO(result.data))
    for shape in deck.slides[0].shapes:
        assert shape.left + shape.width <= deck.slide_width
        assert shape.top + shape.height <= deck.slide_height


def test_long_bullet_gets_finding():
    result = generate("pptx", {"slides": [{"title": "T", "bullets": ["word " * 50]}]})
    assert any(f["code"] == "long_bullet" for f in result.findings)
    assert result.validation_state == "passed_with_warnings"


def test_bounds_reject_excess_slides_bullets_and_title():
    with pytest.raises(DocumentGenerationError, match="60 slides"):
        generate("pptx", {"slides": [{"title": f"S{i}"} for i in range(61)]})
    with pytest.raises(DocumentGenerationError, match="12 bullets"):
        generate("pptx", {"slides": [{"title": "T", "bullets": [f"b{i}" for i in range(13)]}]})
    with pytest.raises(DocumentGenerationError, match="200 characters"):
        generate("pptx", {"slides": [{"title": "x" * 201}]})


def test_bounds_accept_the_limits_exactly():
    result = generate("pptx", {"slides": [{"title": "x" * 200, "bullets": [f"b{i}" for i in range(12)]}
                                          for _ in range(60)]})
    assert result.page_count == 60
    assert len(Presentation(io.BytesIO(result.data)).slides) == 60


@pytest.mark.parametrize("spec,match", [
    ({}, "slides"),
    ({"slides": "no"}, "must be"),
    ({"slides": []}, "at least one"),
    ({"slides": ["nope"]}, "must be an object"),
    ({"slides": [{"bullets": []}]}, "title must be a string"),
    ({"slides": [{"title": 7}]}, "title must be a string"),
    ({"slides": [{"title": "T", "layout": "comic"}]}, "layout must be one of"),
    ({"slides": [{"title": "T", "subtitle": 3}]}, "subtitle must be a string"),
    ({"slides": [{"title": "T", "bullets": "three"}]}, "bullets must be an array"),
    ({"slides": [{"title": "T", "bullets": [1]}]}, "bullet 1 must be a string"),
    ({"slides": [{"title": "T", "notes": 5}]}, "notes must be a string"),
    ({"slides": [{"title": "T", "table": "wide"}]}, "table must be an object"),
    ({"slides": [{"title": "T", "table": {"header": "h", "rows": []}}]}, "non-empty array"),
    ({"slides": [{"title": "T", "table": {"header": ["a"], "rows": [["a", "b"]]}}]}, "cells like the header"),
    ({"slides": [{"title": "T", "table": {"header": ["a"], "rows": [[1]]}}]}, "cell must be a string"),
    ({"slides": [{"title": "T\tX"}]}, "tabs or control characters"),
    ({"slides": [{"title": "T", "bullets": ["a\x0bb"]}]}, "tabs or control characters"),
    ({"slides": [{"title": "T", "notes": "a\x1fb"}]}, "tabs or control characters"),
])
def test_bad_specs_raise_instead_of_coercing(spec, match):
    with pytest.raises(DocumentGenerationError, match=match):
        generate("pptx", spec)


def test_notes_normalize_newlines_and_roundtrip():
    result = generate("pptx", {"slides": [{"title": "T", "notes": "a\r\nb\rc"}]})
    deck = Presentation(io.BytesIO(result.data))
    assert deck.slides[0].notes_slide.notes_text_frame.text == "a\nb\nc"


def test_reopen_gate_rejects_out_of_bounds_geometry(monkeypatch):
    """If placement ever escapes the sheet, the parse-back gate must refuse the bytes."""
    monkeypatch.setattr(pptx_gen, "_BODY_BOTTOM", 8.6)  # past the 7.5" sheet
    with pytest.raises(DocumentGenerationError, match="overflows the slide sheet"):
        generate("pptx", {"slides": [{"title": "T", "bullets": ["b"]}]})
