"""DOCX generator: structure round-trip, bounds, and the E7/theme-indirection defenses."""

from __future__ import annotations

import io

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from hutch.documents.docgen import DocumentGenerationError, generate, supported_formats
from hutch.documents.docgen import docxgen

BASE_FONT = "Liberation Sans"
_THEME_ATTR_NAMES = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")


def _full_spec() -> dict:
    return {"title": "Quarterly Report", "blocks": [
        {"h1": "Quarterly Report"},
        {"h2": "Summary"},
        {"h3": "Detail"},
        {"h4": "Fine print"},
        {"p": "Ordinary paragraph"},
        {"p": "Centered", "align": "center"},
        {"bullets": ["first", "second"]},
        {"numbers": ["one", "two"]},
        {"quote": "Something intense"},
        {"caption": "Figure 1"},
        {"table": {"header": ["id", "name"], "rows": [["1", "ada"], ["2", "grace"]]}},
        {"pagebreak": True},
        {"p": "After break"},
    ]}


_FULL_PLAN = [
    ("Heading 1", "Quarterly Report"),
    ("Heading 2", "Summary"),
    ("Heading 3", "Detail"),
    ("Heading 4", "Fine print"),
    ("Normal", "Ordinary paragraph"),
    ("Normal", "Centered"),
    ("List Bullet", "first"),
    ("List Bullet", "second"),
    ("List Number", "one"),
    ("List Number", "two"),
    ("Intense Quote", "Something intense"),
    ("Caption", "Figure 1"),
    ("Normal", ""),
    ("Normal", "After break"),
]


def test_docx_is_a_registered_format():
    assert "docx" in supported_formats()


def test_full_document_roundtrip_matches_spec():
    result = generate("docx", _full_spec())
    assert result.validation_state == "passed" and result.findings == []
    assert result.preview_mode == "structure"
    reopened = docx.Document(io.BytesIO(result.data))
    assert [(p.style.name, p.text) for p in reopened.paragraphs] == _FULL_PLAN
    assert len(reopened.tables) == 1
    table = reopened.tables[0]
    assert (len(table.rows), len(table.columns)) == (3, 2)
    assert [cell.text for cell in table.rows[0].cells] == ["id", "name"]
    assert [cell.text for cell in table.rows[1].cells] == ["1", "ada"]
    assert table.style.name == "Table Grid"
    centered = reopened.paragraphs[5]
    assert centered.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert reopened.paragraphs[4].alignment is None
    assert reopened.paragraphs[12].paragraph_format.page_break_before is True


def test_core_properties_title_author_language():
    titled = generate("docx", _full_spec())
    props = docx.Document(io.BytesIO(titled.data)).core_properties
    assert props.title == "Quarterly Report"
    assert props.author == "hutch"
    assert props.language == "en-US"
    default = docx.Document(io.BytesIO(
        generate("docx", {"blocks": [{"p": "x"}]}).data)).core_properties
    assert default.title == docxgen.DEFAULT_TITLE
    assert default.author == "hutch" and default.language == "en-US"


def test_base_style_fonts_direct_and_theme_free():
    """python-docx#1089: the written styles carry the explicit font, never theme indirection."""
    result = generate("docx", {"blocks": [{"h1": "T"}]})
    styles = docx.Document(io.BytesIO(result.data)).styles
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        rfonts = styles[name].element.rPr.rFonts
        for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
            assert rfonts.get(qn(f"w:{attr}")) == BASE_FONT, name
        for attr in _THEME_ATTR_NAMES:
            assert rfonts.get(qn(f"w:{attr}")) is None, \
                f"{name} must not depend on theme inheritance"


def test_table_roundtrip_dimensions_and_repeated_header():
    """E7: tables survive the round-trip with intact geometry; the first row repeats."""
    rows = [[str(i), f"item-{i}"] for i in range(4)]
    result = generate("docx", {"blocks": [
        {"h1": "Inventory"},
        {"table": {"header": ["no", "item"], "rows": rows}}]})
    reopened = docx.Document(io.BytesIO(result.data))
    assert len(reopened.tables) == 1
    table = reopened.tables[0]
    assert (len(table.rows), len(table.columns)) == (5, 2)
    assert [cell.text for cell in table.rows[0].cells] == ["no", "item"]
    for index, row in enumerate(rows, start=1):
        assert [cell.text for cell in table.rows[index].cells] == row
    first_tr_pr = table.rows[0]._tr.trPr
    assert first_tr_pr is not None and first_tr_pr.find(qn("w:tblHeader")) is not None
    other_tr_pr = table.rows[2]._tr.trPr
    assert other_tr_pr is None or other_tr_pr.find(qn("w:tblHeader")) is None


def test_generator_string_names_library_and_version():
    result = generate("docx", {"blocks": [{"p": "hi"}]})
    assert result.generator == f"hutch docgen docx 1.0 (python-docx {docx.__version__})"


def test_identical_specs_produce_identical_bytes():
    spec = _full_spec()
    assert generate("docx", spec).data == generate("docx", spec).data


def test_validation_gate_rejects_unreopenable_and_empty_bytes():
    with pytest.raises(DocumentGenerationError, match="did not reopen"):
        docxgen._validate_structure(b"not a docx container", [], [])
    buffer = io.BytesIO()
    docx.Document().save(buffer)
    with pytest.raises(DocumentGenerationError, match="empty"):
        docxgen._validate_structure(buffer.getvalue(), [], [])


def test_preview_outline_shape_and_bound():
    result = generate("docx", _full_spec())
    assert result.preview_data == {"kind": "outline", "items": [
        {"level": 1, "text": "Quarterly Report"},
        {"level": 2, "text": "Summary"},
        {"level": 3, "text": "Detail"},
        {"level": 4, "text": "Fine print"},
    ]}
    many = generate("docx", {"blocks": [{"h1": f"Section {i}"} for i in range(250)]})
    items = many.preview_data["items"]
    assert len(items) == docxgen.MAX_PREVIEW_ITEMS == 200
    assert items[0] == {"level": 1, "text": "Section 0"}
    assert items[-1] == {"level": 1, "text": "Section 199"}


def test_table_over_50_rows_warns_as_page_split():
    def spec(rows: int) -> dict:
        return {"blocks": [{"table": {"header": ["a"], "rows": [[str(i)] for i in range(rows)]}}]}

    warned = generate("docx", spec(51))
    assert warned.validation_state == "passed_with_warnings"
    assert [f["code"] for f in warned.findings] == ["docx_table_page_split"]
    quiet = generate("docx", spec(50))
    assert quiet.validation_state == "passed" and quiet.findings == []


def test_bounds_are_refused():
    with pytest.raises(DocumentGenerationError, match="2,000 blocks"):
        generate("docx", {"blocks": [{"p": "b"}] * 2_001})
    with pytest.raises(DocumentGenerationError, match="exceeds 500 rows"):
        generate("docx", {"blocks": [{"table": {"header": ["a"], "rows": [["x"]] * 501}}]})
    with pytest.raises(DocumentGenerationError, match="exceeds 100,000 characters"):
        generate("docx", {"blocks": [{"p": "x" * 100_001}]})
    with pytest.raises(DocumentGenerationError, match="exceeds 100,000 characters"):
        generate("docx", {"blocks": [{"table": {"header": ["y" * 100_001], "rows": []}}]})


@pytest.mark.parametrize("spec,match", [
    ({"text": "not the right key"}, "unknown docx spec key"),
    ({"blocks": "nope"}, "blocks"),
    ({"blocks": []}, "at least one block"),
    ({"blocks": ["not-an-object"]}, "must be an object"),
    ({"blocks": [{"h5": "no such heading"}]}, "unknown key"),
    ({"blocks": [{"h1": "a", "p": "b"}]}, "exactly one"),
    ({"blocks": [{"h1": "a", "align": "left"}]}, "only valid with p"),
    ({"blocks": [{"p": "a", "align": "middle"}]}, "align must be one of"),
    ({"blocks": [{"p": "a", "align": None}]}, "align must be one of"),
    ({"blocks": [{"bullets": "nope"}]}, "array of strings"),
    ({"blocks": [{"bullets": [1]}]}, "must be a string"),
    ({"blocks": [{"numbers": [None]}]}, "must be a string"),
    ({"blocks": [{"table": {"rows": []}}]}, "table"),
    ({"blocks": [{"table": {"header": ["a"], "rows": "nope"}}]}, "array"),
    ({"blocks": [{"table": {"header": [], "rows": []}}]}, "at least one column"),
    ({"blocks": [{"table": {"header": ["a"], "rows": [42]}}]}, "array of strings"),
    ({"blocks": [{"table": {"header": ["a"], "rows": [[1]]}}]}, "must be a string"),
    ({"blocks": [{"table": {"header": ["a"], "rows": [["b"], ["c", "d"]]}}]}, "cells, expected 1"),
    ({"blocks": [{"table": {"header": ["a"], "rows": [["x"]] * 501}}]}, "500 rows"),
    ({"blocks": [{"pagebreak": False}]}, "pagebreak"),
    ({"blocks": [{"pagebreak": "yes"}]}, "pagebreak"),
    ({"blocks": [{"p": "bad\x00char"}]}, "not valid in OOXML"),
    ({"blocks": [{"p": "ok"}], "extra": 1}, "unknown docx spec key"),
    ({"blocks": [{"p": "ok"}], "title": 3}, "title must be a string"),
])
def test_invalid_specs_raise_without_coercion(spec: dict, match: str):
    with pytest.raises(DocumentGenerationError, match=match):
        generate("docx", spec)


def test_non_object_spec_refused_at_dispatch():
    with pytest.raises(DocumentGenerationError, match="must be an object"):
        generate("docx", ["not", "an", "object"])
