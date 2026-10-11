"""XLSX generator: XlsxWriter writes, openpyxl independently re-opens before promotion.

Constraints this module enforces (failure classes E6/E12 in
``docs/research/document-pipeline-failures-2026-10-10.md``):

- Every formula is written with an explicit cached result (``write_formula(..., value=)``)
  in an auto-calc workbook, because XLSX stores last-computed values and non-Excel readers
  do not recalculate (XlsxWriter#864/#877). A formula cell without a declared value is a
  bad spec, not a zero. After writing, the bytes are reopened with openpyxl
  ``data_only=True`` and every formula cell must expose a cached value equal to the value
  the specification supplied; for the simple aggregates the declared value is additionally
  recomputed from the specification's own source cells, so a spec that misstates its own
  arithmetic fails the gate. Formulas hutch cannot recompute keep the declared cached value
  and downgrade the result to ``passed_with_warnings``.
- Cell types are explicit and verified by the independent reader: strings keep leading
  zeros (IDs survive as text), numbers stay numeric, dates are real datetime cells with a
  ``yyyy-mm-dd`` number format.
- Sheet names are sanitized before the writer sees them: Excel-invalid characters removed,
  31-character cap, case-insensitive deduplication; every change is disclosed as a finding.

Bounds (sheets, rows, columns) are hard limits: exceeding them is an error, never a silent
truncation or coercion.
"""

from __future__ import annotations

import datetime as dt
import io
import math
import re
import zipfile
from dataclasses import dataclass
from importlib.metadata import version as _dist_version
from typing import Any

import xlsxwriter
from openpyxl import load_workbook

from . import DocumentGenerationError, GeneratedDocument, register

GENERATOR = (f"hutch docgen xlsx 1.0 (XlsxWriter {_dist_version('XlsxWriter')}"
             f" + openpyxl {_dist_version('openpyxl')})")

MAX_SHEETS = 20
MAX_SHEET_ROWS = 100_000
MAX_SHEET_COLS = 256
MAX_COL_WIDTH = 255
#: the grid sample the structure preview carries: first N rows of each sheet, header included
PREVIEW_SAMPLE_ROWS = 8
PREVIEW_CELL_CHARS = 120
#: cap on source cells scanned while recomputing formulas, so a spec cannot make the gate
#: quadratic; past it, remaining formulas are disclosed as unrecomputed warnings instead
MAX_RECOMPUTE_CELLS = 2_000_000

#: aggregates the gate recomputes from the specification's own source cells (E6)
_RECOMPUTABLE = frozenset({"SUM", "AVERAGE", "COUNT", "MIN", "MAX"})
_SHEET_NAME_BAD_CHARS = re.compile(r"[\[\]:*?/\\]")
_FORMULA_HEAD = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
_SIMPLE_FORMULA = re.compile(r"^\s*([A-Za-z_][A-Za-z0-9_.]*)\s*\((.*)\)\s*$", re.S)
_ARGUMENT = re.compile(
    r"^\s*(\$?[A-Za-z]{1,3}\$?[0-9]{1,7}(?::\$?[A-Za-z]{1,3}\$?[0-9]{1,7})?|-?[0-9]+(?:\.[0-9]+)?)\s*$")
_CELL_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?([0-9]{1,7})$")
_ISO_DATE = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$")
_SHEET_KEYS = frozenset({"name", "header", "rows", "freeze_header", "col_widths"})
_FORMULA_KEYS = frozenset({"formula", "value"})

Cell = tuple[str, Any]  # ("str", text) | ("num", number) | ("date", dt.date) | ("formula", body, value)


class _UnsupportedFormula(ValueError):
    """The formula is not one of the simple aggregates the gate can recompute."""

    def __init__(self, function: str, *, budget_exhausted: bool = False) -> None:
        super().__init__(function)
        self.function = function
        self.budget_exhausted = budget_exhausted


@dataclass
class _Sheet:
    requested_name: str
    name: str
    header: list[str]
    rows: list[list[Cell]]
    freeze: bool
    col_widths: list[float]
    width: int


class _Budget:
    """Remaining source cells the recompute pass may scan for this document.

    The limit is read at construction time so the module constant governs every instance.
    """

    def __init__(self) -> None:
        self.remaining = MAX_RECOMPUTE_CELLS

    def spend(self, cells: int) -> bool:
        if cells > self.remaining:
            return False
        self.remaining -= cells
        return True


def _bad_spec(detail: str) -> DocumentGenerationError:
    return DocumentGenerationError(f"bad xlsx specification: {detail}")


def _sheet_name(raw: str, index: int, used: set[str]) -> str:
    """Sanitize one sheet name; Excel names are <=31 chars, without []:*?/\\ and unique."""
    name = _SHEET_NAME_BAD_CHARS.sub("", raw).strip("'")
    if not name:
        name = f"Sheet{index + 1}"
    if len(name) > 31:
        name = name[:31]
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f"_{n}"
        candidate = name[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _column_index(letters: str) -> int:
    index = 0
    for ch in letters.upper():
        index = index * 26 + (ord(ch) - ord("A") + 1)
    return index - 1


def _address(row0: int, col0: int) -> str:
    """A1-style address of a 0-based grid position (header included in the row number)."""
    letters, c = "", col0 + 1
    while c:
        c, rem = divmod(c - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return f"{letters}{row0 + 1}"


def _excel_serial(day: dt.date) -> float:
    """Days since the 1900 date system epoch; dates are numbers inside a spreadsheet."""
    return float((day - dt.date(1899, 12, 30)).days)


def _close(left: float, right: float) -> bool:
    return math.isclose(left, right, rel_tol=1e-9, abs_tol=1e-9)


def _normalize_cell(cell: Any, where: str) -> Cell:
    if isinstance(cell, str):
        return ("str", cell)
    if isinstance(cell, bool) or cell is None:
        raise _bad_spec(f"{where}: cells must be strings, numbers, or {{formula|date}} objects")
    if isinstance(cell, (int, float)):
        if not math.isfinite(cell):
            raise _bad_spec(f"{where}: numbers must be finite")
        return ("num", cell)
    if not isinstance(cell, dict):
        raise _bad_spec(f"{where}: cells must be strings, numbers, or {{formula|date}} objects")
    keys = set(cell)
    if keys == {"formula", "value"}:
        formula, value = cell["formula"], cell["value"]
        if not isinstance(formula, str) or not formula.strip():
            raise _bad_spec(f"{where}: 'formula' must be a non-empty string")
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise _bad_spec(f"{where}: a formula cell's cached 'value' must be a finite number")
        return ("formula", formula.lstrip("=").strip(), value)
    if keys == {"date"}:
        text = cell["date"]
        if not isinstance(text, str) or not _ISO_DATE.fullmatch(text):
            raise _bad_spec(f"{where}: 'date' must be a 'YYYY-MM-DD' string")
        try:
            return ("date", dt.date.fromisoformat(text))
        except ValueError:
            raise _bad_spec(f"{where}: {text!r} is not a real calendar date") from None
    raise _bad_spec(f"{where}: cell objects must be exactly {{'formula', 'value'}} or {{'date'}}")


def _normalize_sheets(spec: dict[str, Any]) -> tuple[list[_Sheet], list[dict[str, str]]]:
    sheets_spec = spec.get("sheets")
    if not isinstance(sheets_spec, list) or not sheets_spec:
        raise _bad_spec('content must be {"sheets": [{"name", "header"?, "rows", ...}, ...]} and non-empty')
    if len(sheets_spec) > MAX_SHEETS:
        raise _bad_spec(f"a workbook may hold at most {MAX_SHEETS} sheets")
    sheets: list[_Sheet] = []
    findings: list[dict[str, str]] = []
    used_names: set[str] = set()
    for index, sheet_spec in enumerate(sheets_spec):
        if not isinstance(sheet_spec, dict):
            raise _bad_spec(f"sheet {index} must be an object")
        unknown = set(sheet_spec) - _SHEET_KEYS
        if unknown:
            raise _bad_spec(f"sheet {index} has unsupported keys: {', '.join(sorted(unknown))}")
        requested = sheet_spec.get("name")
        if not isinstance(requested, str):
            raise _bad_spec(f"sheet {index}: 'name' is required and must be a string")
        name = _sheet_name(requested, index, used_names)
        if name != requested:
            findings.append({"code": "sheet_name_sanitized",
                             "detail": f"sheet name {requested!r} was written as {name!r}"})
        header_spec = sheet_spec.get("header")
        if header_spec is not None:
            if not isinstance(header_spec, list) or any(not isinstance(h, str) for h in header_spec):
                raise _bad_spec(f"sheet {index}: 'header' must be a list of strings")
        rows_spec = sheet_spec.get("rows")
        if not isinstance(rows_spec, list):
            raise _bad_spec(f"sheet {index}: 'rows' must be a list of rows")
        if len(rows_spec) > MAX_SHEET_ROWS:
            raise _bad_spec(f"sheet {index} exceeds {MAX_SHEET_ROWS:,} rows")
        header: list[str] = header_spec or []
        offset = 1 if header else 0
        if len(header) > MAX_SHEET_COLS:
            raise _bad_spec(f"sheet {index} exceeds {MAX_SHEET_COLS} columns")
        rows: list[list[Cell]] = []
        width = len(header)
        for r, row_spec in enumerate(rows_spec):
            if not isinstance(row_spec, list):
                raise _bad_spec(f"sheet {index} row {r + 1} must be an array")
            if len(row_spec) > MAX_SHEET_COLS:
                raise _bad_spec(f"sheet {index} exceeds {MAX_SHEET_COLS} columns")
            rows.append([_normalize_cell(cell, f"sheet {index} row {r + 1} column {c + 1}")
                         for c, cell in enumerate(row_spec)])
            width = max(width, len(row_spec))
        freeze_spec = sheet_spec.get("freeze_header", False)
        if not isinstance(freeze_spec, bool):
            raise _bad_spec(f"sheet {index}: 'freeze_header' must be a boolean")
        if freeze_spec and not header:
            raise _bad_spec(f"sheet {index}: 'freeze_header' requires a 'header'")
        widths_spec = sheet_spec.get("col_widths", [])
        if not isinstance(widths_spec, list) or any(
                isinstance(w, bool) or not isinstance(w, (int, float)) for w in widths_spec):
            raise _bad_spec(f"sheet {index}: 'col_widths' must be a list of numbers")
        if any(w < 0 or w > MAX_COL_WIDTH for w in widths_spec):
            raise _bad_spec(f"sheet {index}: column widths must be between 0 and {MAX_COL_WIDTH}")
        if len(widths_spec) > width:
            raise _bad_spec(f"sheet {index}: 'col_widths' covers {len(widths_spec)} columns"
                            f" but the sheet has {width}")
        sheets.append(_Sheet(requested_name=requested, name=name, header=header, rows=rows,
                             freeze=freeze_spec, col_widths=[float(w) for w in widths_spec],
                             width=width))
    return sheets, findings


def _numeric_values(sheet: _Sheet) -> dict[tuple[int, int], float]:
    """Numeric value at each 0-based grid position (header row included), for recompute."""
    values: dict[tuple[int, int], float] = {}
    offset = 1 if sheet.header else 0
    for r, row in enumerate(sheet.rows):
        for c, cell in enumerate(row):
            if cell[0] == "num":
                values[(r + offset, c)] = float(cell[1])
            elif cell[0] == "date":
                values[(r + offset, c)] = _excel_serial(cell[1])
            elif cell[0] == "formula":
                values[(r + offset, c)] = float(cell[2])
    return values


def _argument_values(token: str, values: dict[tuple[int, int], float], budget: _Budget) -> list[float]:
    if ":" in token:
        corners = []
        for bound in token.split(":"):
            match = _CELL_REF.fullmatch(bound)
            assert match is not None
            corners.append((_column_index(match.group(1)), int(match.group(2)) - 1))
        (col1, row1), (col2, row2) = corners
        row_start, row_end = sorted((row1, row2))
        col_start, col_end = sorted((col1, col2))
        if not budget.spend((row_end - row_start + 1) * (col_end - col_start + 1)):
            raise _UnsupportedFormula("", budget_exhausted=True)
        return [values[(row, col)]
                for row in range(row_start, row_end + 1)
                for col in range(col_start, col_end + 1)
                if (row, col) in values]
    match = _CELL_REF.fullmatch(token)
    assert match is not None
    if not budget.spend(1):
        raise _UnsupportedFormula("", budget_exhausted=True)
    number = values.get((int(match.group(2)) - 1, _column_index(match.group(1))))
    return [number] if number is not None else []


def _recompute(body: str, values: dict[tuple[int, int], float], budget: _Budget) -> float:
    """Recompute a simple aggregate from the specification's source cells (E6)."""
    parsed = _SIMPLE_FORMULA.fullmatch(body)
    if parsed is None:
        raise _UnsupportedFormula(body)
    function, arguments = parsed.group(1).upper(), parsed.group(2)
    if function not in _RECOMPUTABLE:
        raise _UnsupportedFormula(function)
    numbers: list[float] = []
    for argument in arguments.split(","):
        match = _ARGUMENT.fullmatch(argument)
        if match is None:
            raise _UnsupportedFormula(function)
        token = match.group(1)
        if ":" in token or _CELL_REF.fullmatch(token):
            numbers.extend(_argument_values(token, values, budget))
        else:
            numbers.append(float(token))
    if function == "COUNT":
        return float(len(numbers))
    if function == "SUM":
        return math.fsum(numbers)
    if function == "AVERAGE":
        if not numbers:
            raise DocumentGenerationError(
                f"bad xlsx specification: {body!r} averages no numeric cells (Excel shows #DIV/0!)")
        return math.fsum(numbers) / len(numbers)
    if function == "MIN":
        return min(numbers) if numbers else 0.0
    return max(numbers) if numbers else 0.0


def _verify_declared_values(sheets: list[_Sheet]) -> list[dict[str, str]]:
    """Compare each recomputable formula's declared value against its own source cells."""
    findings: list[dict[str, str]] = []
    budget = _Budget()
    for sheet in sheets:
        if not any(cell[0] == "formula" for row in sheet.rows for cell in row):
            continue
        values = _numeric_values(sheet)
        offset = 1 if sheet.header else 0
        for r, row in enumerate(sheet.rows):
            for c, cell in enumerate(row):
                if cell[0] != "formula":
                    continue
                body, declared = cell[1], float(cell[2])
                try:
                    computed = _recompute(body, values, budget)
                except _UnsupportedFormula as unsupported:
                    if unsupported.budget_exhausted:
                        findings.append({"code": "formula_not_recomputed", "detail": (
                            f"{sheet.name}!{_address(r + offset, c)}: the recompute budget was"
                            " exhausted; the cached value comes from the specification")})
                    else:
                        findings.append({"code": "formula_not_recomputed", "detail": (
                            f"{sheet.name}!{_address(r + offset, c)}: function"
                            f" {unsupported.function!r} is not one hutch recomputes;"
                            " the cached value comes from the specification")})
                    continue
                if not _close(computed, declared):
                    raise DocumentGenerationError(
                        f"bad xlsx specification: {sheet.name}!{_address(r + offset, c)}"
                        f" declares {declared} but its source cells recompute to {computed}")
    return findings


def _write_workbook(sheets: list[_Sheet]) -> bytes:
    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    workbook.set_calc_mode("auto")
    date_format = workbook.add_format({"num_format": "yyyy-mm-dd"})
    for sheet in sheets:
        try:
            worksheet = workbook.add_worksheet(sheet.name)
        except (xlsxwriter.exceptions.DuplicateWorksheetName,
                xlsxwriter.exceptions.InvalidWorksheetName) as exc:
            raise DocumentGenerationError(f"bad xlsx specification: sheet name: {exc}") from None
        offset = 0
        for c, text in enumerate(sheet.header):
            worksheet.write_string(0, c, text)
            offset = 1
        for r, row in enumerate(sheet.rows):
            for c, cell in enumerate(row):
                excel_row = r + offset
                if cell[0] == "str":
                    worksheet.write_string(excel_row, c, cell[1])
                elif cell[0] == "num":
                    worksheet.write_number(excel_row, c, cell[1])
                elif cell[0] == "date":
                    day = cell[1]
                    worksheet.write_datetime(
                        excel_row, c, dt.datetime(day.year, day.month, day.day), date_format)
                else:
                    worksheet.write_formula(excel_row, c, "=" + cell[1], None, cell[2])
        if sheet.freeze:
            worksheet.freeze_panes(1, 0)
        for c, width in enumerate(sheet.col_widths):
            worksheet.set_column(c, c, width)
    workbook.close()
    return buffer.getvalue()


def _verify_sheet_values(sheet: _Sheet, worksheet: Any) -> None:
    """Read back every written cell through the independent reader and check its type (E12)."""
    expected: list[list[Cell]] = ([[("str", text) for text in sheet.header]] if sheet.header else [])
    expected += sheet.rows
    rows = worksheet.iter_rows(values_only=True)
    for row0, expected_row in enumerate(expected):
        actual_row = next(rows, None)
        if actual_row is None or len(actual_row) < len(expected_row):
            raise DocumentGenerationError(
                f"the written sheet {sheet.name!r} is missing grid row {row0 + 1}")
        for c, cell in enumerate(expected_row):
            actual = actual_row[c]
            where = f"{sheet.name}!{_address(row0, c)}"
            if cell[0] == "str":
                if not isinstance(actual, str) or actual != cell[1]:
                    raise DocumentGenerationError(f"{where}: string cell did not round-trip as text")
            elif cell[0] == "num":
                if isinstance(actual, bool) or not isinstance(actual, (int, float)) or actual != cell[1]:
                    raise DocumentGenerationError(f"{where}: number cell did not round-trip as a number")
            elif cell[0] == "date":
                if not isinstance(actual, dt.datetime) or actual.date() != cell[1]:
                    raise DocumentGenerationError(f"{where}: date cell did not round-trip as a date")
            else:
                declared = float(cell[2])
                if actual is None:
                    raise DocumentGenerationError(
                        f"{where}: formula has no cached value; non-Excel readers would show a blank")
                if isinstance(actual, bool) or not isinstance(actual, (int, float)) \
                        or not _close(float(actual), declared):
                    raise DocumentGenerationError(
                        f"{where}: cached value {actual!r} does not match the declared {declared}")
    if next(rows, None) is not None:
        raise DocumentGenerationError(f"the written sheet {sheet.name!r} has unexpected extra rows")


def _verify_written_workbook(data: bytes, sheets: list[_Sheet]) -> None:
    """The promotion gate: reopen with openpyxl and prove the bytes say what the spec said."""
    cached = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        if cached.sheetnames != [sheet.name for sheet in sheets]:
            raise DocumentGenerationError("the written workbook's sheet names do not match the spec")
        for sheet in sheets:
            _verify_sheet_values(sheet, cached[sheet.name])
    finally:
        cached.close()
    if any(cell[0] == "formula" for sheet in sheets for row in sheet.rows for cell in row) \
            or any(sheet.freeze for sheet in sheets):
        book = load_workbook(io.BytesIO(data), data_only=False)
        try:
            for sheet in sheets:
                worksheet = book[sheet.name]
                if sheet.freeze and worksheet.freeze_panes != "A2":
                    raise DocumentGenerationError(f"sheet {sheet.name!r} did not freeze its header row")
                offset = 1 if sheet.header else 0
                for r, row in enumerate(sheet.rows):
                    for c, cell in enumerate(row):
                        if cell[0] != "formula":
                            continue
                        stored = worksheet.cell(row=r + offset + 1, column=c + 1).value
                        normalized = re.sub(r"[\s$]", "", str(stored)).upper()
                        if normalized != "=" + re.sub(r"[\s$]", "", cell[1]).upper():
                            raise DocumentGenerationError(
                                f"{sheet.name}!{_address(r + offset, c)}: formula did not round-trip")
        finally:
            book.close()


def _preview_cell(cell: Cell) -> str | int | float:
    if cell[0] == "str":
        text = cell[1]
        return text if len(text) <= PREVIEW_CELL_CHARS else text[:PREVIEW_CELL_CHARS] + "…"
    if cell[0] == "num":
        return cell[1]
    if cell[0] == "date":
        return cell[1].isoformat()
    return cell[2]


def _preview(sheets: list[_Sheet]) -> dict[str, Any]:
    previews = []
    for sheet in sheets:
        sample: list[list[str | int | float]] = []
        if sheet.header:
            sample.append([_preview_cell(("str", text)) for text in sheet.header])
        for row in sheet.rows[: PREVIEW_SAMPLE_ROWS - len(sample)]:
            sample.append([_preview_cell(cell) for cell in row])
        previews.append({"name": sheet.name, "rows": len(sheet.rows), "cols": sheet.width,
                         "sample": sample})
    return {"kind": "grid", "sheets": previews}


@register("xlsx")
def _gen_xlsx(spec: dict[str, Any]) -> GeneratedDocument:
    """Render a typed sheet spec and gate the bytes through an independent openpyxl read-back."""
    sheets, name_findings = _normalize_sheets(spec)
    formula_findings = _verify_declared_values(sheets)
    data = _write_workbook(sheets)
    _verify_written_workbook(data, sheets)
    state = "passed_with_warnings" if formula_findings else "passed"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if b'fullCalcOnLoad="1"' not in archive.read("xl/workbook.xml"):
            raise DocumentGenerationError("the written workbook does not request recalculation on open")
    return GeneratedDocument(data=data, generator=GENERATOR, validation_state=state,
                             findings=name_findings + formula_findings,
                             preview_mode="structure", preview_data=_preview(sheets))
