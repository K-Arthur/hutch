"""DOCX generator: typed block specification → WordprocessingML via python-docx.

Content is placed in runs at build time and never reassigned afterwards: assigning
``paragraph.text``/``cell.text`` replaces a paragraph's runs and destroys run-level
formatting (python-docx#1251, failure class E7), so this module has no text assignment
anywhere. Base style fonts are authored directly and the theme indirection attributes
are removed, ``font.name`` alone leaves ``asciiTheme``/``hAnsiTheme`` standing and
OOXML lets a theme attribute supersede the explicit font (python-docx#1089), so output
would otherwise silently depend on the theme of whatever application opened the file.

Generation and validation are one gate: the bytes are reopened from memory and their
paragraph/table structure compared against the specification before they are returned;
a file that will not reopen, differs from the spec, or carries no content raises
``DocumentGenerationError`` and nothing is promoted to storage.
"""

from __future__ import annotations

import io
import re
from typing import Any

import docx
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from . import DocumentGenerationError, GeneratedDocument, register

MAX_BLOCKS = 2_000
MAX_TABLE_ROWS = 500
MAX_TEXT_CHARS = 100_000
MAX_PREVIEW_ITEMS = 200
#: tables longer than this are disclosed as page-splitters (the header row repeats)
TABLE_PAGE_SPLIT_ROWS = 50

BASE_FONT = "Liberation Sans"
AUTHOR = "hutch"
LANGUAGE = "en-US"
DEFAULT_TITLE = "hutch document"

#: characters XML 1.0 forbids in content; rejecting them here fails the spec at the
#: boundary instead of failing the serializer late with a malformed part (E1-adjacent)
_INVALID_XML_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\ud800-\udfff]")

_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4}
_ALIGNMENTS = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
               "right": WD_ALIGN_PARAGRAPH.RIGHT}
_CONTENT_KEYS = frozenset({"h1", "h2", "h3", "h4", "p", "bullets", "numbers", "quote",
                           "table", "caption", "pagebreak"})
_BLOCK_KEYS = _CONTENT_KEYS | {"align"}
#: theme attributes supersede explicit fonts in OOXML; their presence would silently
#: re-introduce the theme inheritance the direct font authoring is meant to remove
_THEME_FONT_ATTRS = (qn("w:asciiTheme"), qn("w:hAnsiTheme"), qn("w:eastAsiaTheme"), qn("w:cstheme"))


def _text_field(value: Any, what: str) -> str:
    """Every string slot of the spec: typed, bounded, and XML-safe, no coercion."""
    if not isinstance(value, str):
        raise DocumentGenerationError(f"{what} must be a string")
    if len(value) > MAX_TEXT_CHARS:
        raise DocumentGenerationError(f"{what} exceeds {MAX_TEXT_CHARS:,} characters")
    if _INVALID_XML_CHARS.search(value) is not None:
        raise DocumentGenerationError(f"{what} contains characters that are not valid in OOXML text")
    return value


def _string_list(value: Any, what: str) -> list[str]:
    if not isinstance(value, list):
        raise DocumentGenerationError(f"{what} must be an array of strings")
    return [_text_field(item, f"{what}[{index}]") for index, item in enumerate(value)]


def _author_base_fonts(document: DocumentObject) -> None:
    """Direct-author ``BASE_FONT`` on Normal and the heading styles, theme attrs removed.

    ``font.name`` writes ``w:ascii``/``w:hAnsi`` but leaves the ``*Theme`` attributes in
    place, and OOXML lets those win (python-docx#1089), deleting them makes the explicit
    font authoritative for western, east-asian, and complex-script runs alike.
    """
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[style_name]
        style.font.name = BASE_FONT
        rfonts = style.element.rPr.rFonts
        for attr in _THEME_FONT_ATTRS:
            rfonts.attrib.pop(attr, None)
        rfonts.set(qn("w:cs"), BASE_FONT)
        rfonts.set(qn("w:eastAsia"), BASE_FONT)


def _add_text_paragraph(document: DocumentObject, style: str, text: str,
                        alignment: WD_ALIGN_PARAGRAPH | None = None) -> Paragraph:
    """Body paragraphs receive runs at build time; ``paragraph.text`` assignment is E7."""
    paragraph = document.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    if alignment is not None:
        paragraph.alignment = alignment
    return paragraph


def _fill_cell(cell: _Cell, text: str) -> None:
    """Cells receive runs, never ``cell.text`` assignment (E7: text assignment replaces runs)."""
    if text:
        cell.paragraphs[0].add_run(text)


def _repeat_header_row(table: "docx.table.Table") -> None:
    """Mark the first row ``w:tblHeader`` so it repeats when the table splits across pages."""
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row_properties.append(marker)


def _validate_structure(data: bytes, expected_paragraphs: list[tuple[str, str]],
                        expected_tables: list[tuple[int, int, tuple[str, ...]]]) -> None:
    """Reopen the written bytes and compare them to the specification before promotion.

    This is the promotion gate (failure class E1): a document that will not reopen,
    whose body differs from the spec, or that has no body content at all is an error,
    never a download.
    """
    try:
        reopened = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentGenerationError(f"the written docx did not reopen: {exc}") from None
    if not reopened.paragraphs and not reopened.tables:
        raise DocumentGenerationError("the written docx is empty")
    actual = [(paragraph.style.name, paragraph.text) for paragraph in reopened.paragraphs]
    if len(actual) != len(expected_paragraphs):
        raise DocumentGenerationError(
            f"the written docx has {len(actual)} body paragraphs, "
            f"expected {len(expected_paragraphs)}")
    for index, (got, expected) in enumerate(zip(actual, expected_paragraphs)):
        if got != expected:
            raise DocumentGenerationError(
                f"the written docx does not match the specification at body paragraph {index}")
    actual_tables = reopened.tables
    if len(actual_tables) != len(expected_tables):
        raise DocumentGenerationError(
            f"the written docx has {len(actual_tables)} tables, expected {len(expected_tables)}")
    for index, (table, (rows, cols, header)) in enumerate(zip(actual_tables, expected_tables)):
        if (len(table.rows), len(table.columns)) != (rows, cols):
            raise DocumentGenerationError(f"table {index} does not have the specified dimensions")
        if tuple(cell.text for cell in table.rows[0].cells) != header:
            raise DocumentGenerationError(f"table {index} header does not match the specification")


@register("docx")
def _gen_docx(spec: dict[str, Any]) -> GeneratedDocument:
    """Render ``{"blocks": [...], "title"?: str}``; validate the bytes before returning them.

    ``title`` is the one optional key beyond the documented shape: it lands in the core
    properties (the tool layer records the display title separately, so the document
    carries a default rather than an empty title).
    """
    unknown = set(spec) - {"blocks", "title"}
    if unknown:
        raise DocumentGenerationError(f"unknown docx spec key(s): {', '.join(sorted(unknown))}")
    blocks = spec.get("blocks")
    if not isinstance(blocks, list):
        raise DocumentGenerationError('docx content must be {"blocks": [...]}')
    if not blocks:
        raise DocumentGenerationError("the document must contain at least one block")
    if len(blocks) > MAX_BLOCKS:
        raise DocumentGenerationError(f"the document exceeds {MAX_BLOCKS:,} blocks")
    title = _text_field(spec["title"], "title") if "title" in spec else DEFAULT_TITLE

    document = docx.Document()
    _author_base_fonts(document)
    properties = document.core_properties
    properties.title = title
    properties.author = AUTHOR
    properties.language = LANGUAGE

    paragraphs: list[tuple[str, str]] = []
    tables: list[tuple[int, int, tuple[str, ...]]] = []
    outline: list[dict[str, Any]] = []
    split_tables: list[int] = []
    for index, block in enumerate(blocks):
        what = f"block {index}"
        if not isinstance(block, dict):
            raise DocumentGenerationError(f"{what} must be an object")
        unknown_keys = set(block) - _BLOCK_KEYS
        if unknown_keys:
            raise DocumentGenerationError(
                f"{what} has unknown key(s): {', '.join(sorted(unknown_keys))}")
        content_keys = set(block) & _CONTENT_KEYS
        if len(content_keys) != 1:
            raise DocumentGenerationError(
                f"{what} must contain exactly one of: {', '.join(sorted(_CONTENT_KEYS))}")
        key = content_keys.pop()
        if "align" in block and key != "p":
            raise DocumentGenerationError(f"{what} align is only valid with p")

        if key in _HEADINGS:
            level = _HEADINGS[key]
            text = _text_field(block[key], f"{what} text")
            _add_text_paragraph(document, f"Heading {level}", text)
            paragraphs.append((f"Heading {level}", text))
            outline.append({"level": level, "text": text})
        elif key == "p":
            text = _text_field(block[key], f"{what} text")
            alignment = None
            if "align" in block:
                align = block["align"]
                if not isinstance(align, str) or align not in _ALIGNMENTS:
                    raise DocumentGenerationError(
                        f"{what} align must be one of: {', '.join(sorted(_ALIGNMENTS))}")
                alignment = _ALIGNMENTS[align]
            _add_text_paragraph(document, "Normal", text, alignment)
            paragraphs.append(("Normal", text))
        elif key in ("bullets", "numbers"):
            style = "List Bullet" if key == "bullets" else "List Number"
            for item in _string_list(block[key], f"{what} {key}"):
                _add_text_paragraph(document, style, item)
                paragraphs.append((style, item))
        elif key in ("quote", "caption"):
            style = "Intense Quote" if key == "quote" else "Caption"
            text = _text_field(block[key], f"{what} text")
            _add_text_paragraph(document, style, text)
            paragraphs.append((style, text))
        elif key == "pagebreak":
            if block[key] is not True:
                raise DocumentGenerationError(f'{what} must be {{"pagebreak": true}}')
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.page_break_before = True
            paragraphs.append(("Normal", ""))
        elif key == "table":
            table_spec = block[key]
            if (not isinstance(table_spec, dict) or "header" not in table_spec
                    or "rows" not in table_spec):
                raise DocumentGenerationError(
                    f'{what} must be {{"table": {{"header": [...], "rows": [[...], ...]}}}}')
            unknown_table_keys = set(table_spec) - {"header", "rows"}
            if unknown_table_keys:
                raise DocumentGenerationError(
                    f"{what} table has unknown key(s): {', '.join(sorted(unknown_table_keys))}")
            header = _string_list(table_spec["header"], f"{what} table header")
            if not header:
                raise DocumentGenerationError(
                    f"{what} table header must have at least one column")
            raw_rows = table_spec["rows"]
            if not isinstance(raw_rows, list):
                raise DocumentGenerationError(f"{what} table rows must be an array")
            if len(raw_rows) > MAX_TABLE_ROWS:
                raise DocumentGenerationError(
                    f"{what} table exceeds {MAX_TABLE_ROWS:,} rows")
            data_rows = []
            for row_index, row in enumerate(raw_rows):
                cells = _string_list(row, f"{what} table row {row_index}")
                if len(cells) != len(header):
                    raise DocumentGenerationError(
                        f"{what} table row {row_index} has {len(cells)} cells, "
                        f"expected {len(header)}")
                data_rows.append(cells)
            table = document.add_table(rows=len(data_rows) + 1, cols=len(header))
            table.style = "Table Grid"
            for column, cell_text in enumerate(header):
                _fill_cell(table.cell(0, column), cell_text)
            for row_index, row in enumerate(data_rows, start=1):
                for column, cell_text in enumerate(row):
                    _fill_cell(table.cell(row_index, column), cell_text)
            _repeat_header_row(table)
            tables.append((len(data_rows) + 1, len(header), tuple(header)))
            if len(data_rows) > TABLE_PAGE_SPLIT_ROWS:
                split_tables.append(len(tables))

    buffer = io.BytesIO()
    try:
        document.save(buffer)
    except ValueError as exc:
        raise DocumentGenerationError(f"the docx library refused the content: {exc}") from None
    from .zippack import normalize_ooxml_zip
    data = normalize_ooxml_zip(buffer.getvalue())
    _validate_structure(data, paragraphs, tables)
    findings = [{"code": "docx_table_page_split",
                 "detail": f"table {position} has more than {TABLE_PAGE_SPLIT_ROWS} rows "
                           "and will split across pages (its header row repeats)"}
                for position in split_tables]
    return GeneratedDocument(
        data=data,
        generator=f"hutch docgen docx 1.0 (python-docx {docx.__version__})",
        validation_state="passed_with_warnings" if findings else "passed",
        findings=findings,
        preview_mode="structure",
        preview_data={"kind": "outline", "items": outline[:MAX_PREVIEW_ITEMS]})
