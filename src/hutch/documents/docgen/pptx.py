"""PPTX deck generator: explicit textboxes, probed-font text fitting, geometric gate.

python-pptx cannot know whether text overflows a slide — autofit is renderer metadata,
not layout (python-pptx#973, closed "nothing to be done") — so
``MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE`` is never written. Instead every text frame is fitted
with ``fit_text(font_file=<probed TTF>)`` so real font metrics bake explicit run sizes
into the stored bytes; when no probed font exists (or fitting is impossible, e.g. a word
too wide to wrap at any size, which crashes ``TextFitter._break_line``) conservative
fixed sizes are applied and the degradation is disclosed as a finding. The saved deck is
re-opened before it may be promoted to assert the slide count, every slide title, and
that every shape lies inside the 16:9 sheet — the overflow check the library lacks.
"""

from __future__ import annotations

import io
import os
import re
from typing import Any

import pptx
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import DocumentGenerationError, GeneratedDocument, register

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MAX_SLIDES = 60
MAX_BULLETS_PER_SLIDE = 12
MAX_TITLE_CHARS = 200
#: bullets longer than this cannot be trusted to stay legible at any fit size; they are
#: disclosed, never truncated
LONG_BULLET_CHARS = 200

#: probed real TTFs, first existing wins — fit_text must compute from true metrics
_FONT_CANDIDATES: tuple[tuple[str, str], ...] = (
    ("/usr/share/fonts/liberation/LiberationSans-Regular.ttf", "Liberation Sans"),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf", "DejaVu Sans"),
)

_TITLE_FIT_MAX_PT = 40
_SUBTITLE_FIT_MAX_PT = 20
_BULLET_FIT_MAX_PT = 18
#: conservative explicit sizes when no probed font exists or fitting is impossible
_TITLE_FALLBACK_PT = 32
_SUBTITLE_FALLBACK_PT = 20
_BULLET_FALLBACK_PT = 16
_TABLE_PT = 12

#: shapes carry these names so the reopen gate can find them without placeholder logic
_TITLE_NAME = "hutch-title"
_SUBTITLE_NAME = "hutch-subtitle"
_BULLETS_NAME = "hutch-bullets"
_TABLE_NAME = "hutch-table"

#: sheet-inch placement rects (left, top, width, height); the reopen gate, not these
#: constants, is what guarantees every saved shape stays inside 13.333 x 7.5
_TITLE_RECT = (0.6, 0.4, 12.13, 1.0)
_SUBTITLE_RECT = (0.6, 1.5, 12.13, 0.5)
_BODY_TOP = 1.6
_BODY_TOP_WITH_SUBTITLE = 2.2
_BODY_BOTTOM = 7.1
_SECTION_TITLE_RECT = (0.6, 2.6, 12.13, 1.4)
_SECTION_SUBTITLE_RECT = (0.6, 4.2, 12.13, 0.6)
_SECTION_BODY_TOP = 4.4
_SECTION_BODY_TOP_WITH_SUBTITLE = 5.0
_TITLE_SLIDE_BODY_RECT = (1.0, 6.2, 11.33, 1.1)

_LAYOUTS = ("title", "title_content", "section")
_BLANK_LAYOUT_INDEX = 6
_TITLE_LAYOUT_INDEX = 0
_HANGING_INDENT_EMU = 228600  # 0.25"

#: python-pptx's text setters split ``\n`` across paragraphs and mangle ``\r``/``\t``
#: into escaped forms, and OOXML cannot represent C0 controls at all — spec text is
#: therefore single-line, or (notes only) ``\n``-separated
_FORBIDDEN_LINE_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\r\t]")
_FORBIDDEN_NOTES_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\t]")


def probe_font() -> tuple[str, str] | None:
    """First existing ``(path, family)`` candidate, or None when none is installed."""
    return next(((path, family) for path, family in _FONT_CANDIDATES if os.path.exists(path)),
                None)


def _line_text(value: Any, what: str) -> str:
    """A string the format can store verbatim: no tabs or control characters."""
    if not isinstance(value, str):
        raise DocumentGenerationError(f"{what} must be a string")
    if _FORBIDDEN_LINE_CHARS.search(value):
        raise DocumentGenerationError(f"{what} must not contain tabs or control characters")
    return value


def _optional_str(item: dict[str, Any], key: str, what: str) -> str | None:
    value = item.get(key)
    return None if value is None else _line_text(value, what)


def _notes_text(item: dict[str, Any], no: int) -> str | None:
    value = item.get("notes")
    if value is None:
        return None
    if not isinstance(value, str):
        raise DocumentGenerationError(f"slide {no} notes must be a string")
    notes = value.replace("\r\n", "\n").replace("\r", "\n")
    if _FORBIDDEN_NOTES_CHARS.search(notes):
        raise DocumentGenerationError(f"slide {no} notes must not contain tabs or control characters")
    return notes


def _table_spec(item: dict[str, Any], no: int) -> tuple[list[str], list[list[str]]] | None:
    value = item.get("table")
    if value is None:
        return None
    if not isinstance(value, dict):
        raise DocumentGenerationError(f"slide {no} table must be an object")
    raw_header = value.get("header")
    raw_rows = value.get("rows")
    if not isinstance(raw_header, list) or not raw_header:
        raise DocumentGenerationError(f"slide {no} table header must be a non-empty array")
    if not isinstance(raw_rows, list):
        raise DocumentGenerationError(f"slide {no} table rows must be an array")
    header = [_line_text(cell, f"slide {no} table header cell") for cell in raw_header]
    rows = []
    for r, raw_row in enumerate(raw_rows, 1):
        if not isinstance(raw_row, list):
            raise DocumentGenerationError(f"slide {no} table row {r} must be an array")
        if len(raw_row) != len(header):
            raise DocumentGenerationError(
                f"slide {no} table row {r} must have exactly {len(header)} cells like the header")
        rows.append([_line_text(cell, f"slide {no} table row {r} cell") for cell in raw_row])
    return header, rows


def _validated_slides(spec: dict[str, Any]) -> list[dict[str, Any]]:
    raw_slides = spec.get("slides")
    if not isinstance(raw_slides, list) or not raw_slides:
        raise DocumentGenerationError(
            'pptx content must be {"slides": [{...}, ...]} with at least one slide')
    if len(raw_slides) > MAX_SLIDES:
        raise DocumentGenerationError(f"pptx exceeds {MAX_SLIDES} slides")
    slides: list[dict[str, Any]] = []
    for no, raw in enumerate(raw_slides, 1):
        if not isinstance(raw, dict):
            raise DocumentGenerationError(f"slide {no} must be an object")
        title = _line_text(raw.get("title"), f"slide {no} title")
        if len(title) > MAX_TITLE_CHARS:
            raise DocumentGenerationError(
                f"slide {no} title exceeds {MAX_TITLE_CHARS} characters")
        layout = raw.get("layout", "title_content")
        if layout not in _LAYOUTS:
            raise DocumentGenerationError(
                f"slide {no} layout must be one of: {', '.join(_LAYOUTS)}")
        raw_bullets = raw.get("bullets")
        if raw_bullets is None:
            bullets: list[str] = []
        else:
            if not isinstance(raw_bullets, list):
                raise DocumentGenerationError(f"slide {no} bullets must be an array")
            if len(raw_bullets) > MAX_BULLETS_PER_SLIDE:
                raise DocumentGenerationError(
                    f"slide {no} exceeds {MAX_BULLETS_PER_SLIDE} bullets")
            bullets = [_line_text(b, f"slide {no} bullet {i}") for i, b in enumerate(raw_bullets, 1)]
        slides.append({"title": title, "layout": layout,
                       "subtitle": _optional_str(raw, "subtitle", f"slide {no} subtitle"),
                       "bullets": bullets, "notes": _notes_text(raw, no),
                       "table": _table_spec(raw, no)})
    return slides


def _set_explicit_size(text_frame: Any, size_pt: int, *, bold: bool = False) -> None:
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size_pt)
            run.font.bold = bold


def _fit_frame(text_frame: Any, what: str, *, font: tuple[str, str] | None, max_pt: int,
               fallback_pt: int, bold: bool, findings: list[dict[str, Any]]) -> None:
    """Bake run sizes from true font metrics; never defer sizing to a renderer (E8).

    ``fit_text`` computes with the probed TTF and writes explicit sizes, so the stored
    bytes already fit — no viewer is trusted to shrink anything. When fitting is
    impossible (no probed font, or an unbreakable word wider than the frame at any size,
    which raises from ``TextFitter._break_line``) conservative fixed sizes are applied
    and the degradation is disclosed.
    """
    if font is None:
        _set_explicit_size(text_frame, fallback_pt, bold=bold)
        return
    path, family = font
    try:
        text_frame.fit_text(font_family=family, max_size=max_pt, bold=bold, font_file=path)
    except (TypeError, ValueError, OSError) as exc:
        _set_explicit_size(text_frame, fallback_pt, bold=bold)
        findings.append({"code": "fit_text_fallback",
                         "detail": f"{what}: text could not be auto-fitted ({exc}); "
                                   f"fixed {fallback_pt}pt applied"})


def _add_textbox(slide: Any, rect: tuple[float, float, float, float], name: str) -> Any:
    left, top, width, height = (Inches(v) for v in rect)
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    return box.text_frame


def _centered(text_frame: Any) -> None:
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER


def _fill_bullets(text_frame: Any, bullets: list[str]) -> None:
    text_frame.text = bullets[0]
    for text in bullets[1:]:
        paragraph = text_frame.add_paragraph()
        paragraph.text = text
    for paragraph in text_frame.paragraphs:
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("marL", str(_HANGING_INDENT_EMU))
        pPr.set("indent", str(-_HANGING_INDENT_EMU))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))


def _style_cell(cell: Any, text: str, *, bold: bool) -> None:
    cell.text = text
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(_TABLE_PT)
            run.font.bold = bold


def _split_body(rect: tuple[float, float, float, float], has_bullets: bool,
                has_table: bool) -> tuple[tuple | None, tuple | None]:
    left, top, width, height = rect
    if has_bullets and has_table:
        bullets_h = round(height * 0.45, 3)
        return ((left, top, width, bullets_h),
                (left, round(top + bullets_h + 0.15, 3), width,
                 round(height - bullets_h - 0.15, 3)))
    if has_bullets:
        return rect, None
    if has_table:
        return None, rect
    return None, None


def _remove(shape: Any) -> None:
    shape._element.getparent().remove(shape._element)


def _build_slide(prs: Presentation, slide_spec: dict[str, Any], no: int,
                 font: tuple[str, str] | None, findings: list[dict[str, Any]]) -> None:
    layout = slide_spec["layout"]
    if layout == "title":
        slide = prs.slides.add_slide(prs.slide_layouts[_TITLE_LAYOUT_INDEX])
        title_shape = slide.shapes.title
        if title_shape is None:
            raise DocumentGenerationError("the title-slide layout has no title placeholder")
        title_shape.name = _TITLE_NAME
        title_shape.text_frame.text = slide_spec["title"]
        _fit_frame(title_shape.text_frame, f"slide {no} title", font=font,
                   max_pt=_TITLE_FIT_MAX_PT, fallback_pt=_TITLE_FALLBACK_PT, bold=True,
                   findings=findings)
        subtitle_shape = next((ph for ph in slide.placeholders
                               if ph.placeholder_format.idx == 1), None)
        if slide_spec["subtitle"] is None:
            if subtitle_shape is not None:
                _remove(subtitle_shape)
        else:
            if subtitle_shape is None:
                raise DocumentGenerationError("the title-slide layout has no subtitle placeholder")
            subtitle_shape.name = _SUBTITLE_NAME
            subtitle_shape.text_frame.text = slide_spec["subtitle"]
            _fit_frame(subtitle_shape.text_frame, f"slide {no} subtitle", font=font,
                       max_pt=_SUBTITLE_FIT_MAX_PT, fallback_pt=_SUBTITLE_FALLBACK_PT,
                       bold=False, findings=findings)
        body_rect = _TITLE_SLIDE_BODY_RECT
    else:
        slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT_INDEX])
        title_rect = _SECTION_TITLE_RECT if layout == "section" else _TITLE_RECT
        title_frame = _add_textbox(slide, title_rect, _TITLE_NAME)
        title_frame.text = slide_spec["title"]
        if layout == "section":
            _centered(title_frame)
        _fit_frame(title_frame, f"slide {no} title", font=font,
                   max_pt=_TITLE_FIT_MAX_PT, fallback_pt=_TITLE_FALLBACK_PT, bold=True,
                   findings=findings)
        if slide_spec["subtitle"] is not None:
            sub_rect = (_SECTION_SUBTITLE_RECT if layout == "section" else _SUBTITLE_RECT)
            subtitle_frame = _add_textbox(slide, sub_rect, _SUBTITLE_NAME)
            subtitle_frame.text = slide_spec["subtitle"]
            if layout == "section":
                _centered(subtitle_frame)
            _fit_frame(subtitle_frame, f"slide {no} subtitle", font=font,
                       max_pt=_SUBTITLE_FIT_MAX_PT, fallback_pt=_SUBTITLE_FALLBACK_PT,
                       bold=False, findings=findings)
        body_top = (_SECTION_BODY_TOP_WITH_SUBTITLE if layout == "section" else _BODY_TOP_WITH_SUBTITLE) \
            if slide_spec["subtitle"] is not None \
            else (_SECTION_BODY_TOP if layout == "section" else _BODY_TOP)
        body_rect = (_TITLE_RECT[0], body_top, _TITLE_RECT[2], round(_BODY_BOTTOM - body_top, 3))

    has_bullets, has_table = bool(slide_spec["bullets"]), slide_spec["table"] is not None
    if has_bullets or has_table:
        bullets_rect, table_rect = _split_body(body_rect, has_bullets, has_table)
        if bullets_rect is not None:
            bullets_frame = _add_textbox(slide, bullets_rect, _BULLETS_NAME)
            _fill_bullets(bullets_frame, slide_spec["bullets"])
            _fit_frame(bullets_frame, f"slide {no} bullets", font=font,
                       max_pt=_BULLET_FIT_MAX_PT, fallback_pt=_BULLET_FALLBACK_PT,
                       bold=False, findings=findings)
        if table_rect is not None:
            header, rows = slide_spec["table"]
            left, top, width, height = table_rect
            graphic = slide.shapes.add_table(len(rows) + 1, len(header), Inches(left),
                                             Inches(top), Inches(width), Inches(height))
            graphic.name = _TABLE_NAME
            table = graphic.table
            for c, text in enumerate(header):
                _style_cell(table.cell(0, c), text, bold=True)
            for r, row in enumerate(rows, 1):
                for c, text in enumerate(row):
                    _style_cell(table.cell(r, c), text, bold=False)

    if slide_spec["notes"] is not None:
        slide.notes_slide.notes_text_frame.text = slide_spec["notes"]


def _validate_reopened(data: bytes, slides: list[dict[str, Any]]) -> None:
    """Parse-back of the exact stored bytes: count, titles, and the shape-bounds check
    python-pptx itself does not perform."""
    deck = Presentation(io.BytesIO(data))
    if deck.slide_width != SLIDE_WIDTH or deck.slide_height != SLIDE_HEIGHT:
        raise DocumentGenerationError("the saved deck is not the 16:9 sheet it was built as")
    if len(deck.slides) != len(slides):
        raise DocumentGenerationError(
            f"the saved deck has {len(deck.slides)} slides, expected {len(slides)}")
    for no, (slide, expected) in enumerate(zip(deck.slides, slides), 1):
        named = {shape.name: shape for shape in slide.shapes}
        title_shape = named.get(_TITLE_NAME)
        if (title_shape is None or not title_shape.has_text_frame
                or title_shape.text_frame.text != expected["title"]):
            raise DocumentGenerationError(f"slide {no}: title did not survive the round-trip")
        for shape in slide.shapes:
            if shape.left is None or shape.top is None \
                    or shape.width is None or shape.height is None:
                raise DocumentGenerationError(
                    f"slide {no}: shape {shape.name!r} has no explicit geometry; "
                    "its bounds cannot be verified")
            if shape.left + shape.width > deck.slide_width \
                    or shape.top + shape.height > deck.slide_height:
                raise DocumentGenerationError(
                    f"slide {no}: shape {shape.name!r} overflows the slide sheet")


@register("pptx")
def _gen_pptx(spec: dict[str, Any]) -> GeneratedDocument:
    slides = _validated_slides(spec)
    findings: list[dict[str, Any]] = []
    for no, slide_spec in enumerate(slides, 1):
        for i, bullet in enumerate(slide_spec["bullets"], 1):
            if len(bullet) > LONG_BULLET_CHARS:
                findings.append({"code": "long_bullet",
                                 "detail": f"slide {no} bullet {i} is {len(bullet)} characters "
                                           "and may render small"})
    font = probe_font()
    if font is None:
        findings.append({"code": "fit_text_unavailable",
                         "detail": "no probed TrueType font found on this host; "
                                   "conservative fixed font sizes were applied"})
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    for no, slide_spec in enumerate(slides, 1):
        _build_slide(prs, slide_spec, no, font, findings)
    buffer = io.BytesIO()
    prs.save(buffer)
    data = buffer.getvalue()
    _validate_reopened(data, slides)
    return GeneratedDocument(
        data=data,
        generator=f"hutch docgen pptx 1.0 (python-pptx {pptx.__version__})",
        validation_state="passed_with_warnings" if findings else "passed",
        findings=findings,
        preview_mode="structure",
        preview_data={"kind": "slides", "items": [
            {"title": slide_spec["title"], "bullet_count": len(slide_spec["bullets"]),
             "has_table": slide_spec["table"] is not None,
             "has_notes": slide_spec["notes"] is not None}
            for slide_spec in slides]},
        page_count=len(slides))
